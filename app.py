import os
from pathlib import Path
from flask import Flask, render_template, request
from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent
EXCEL_PATH = Path(os.environ.get("DOCTORS_EXCEL_PATH", BASE_DIR / "doctors.xlsx"))
app = Flask(__name__, template_folder="templates", static_folder="assets", static_url_path="/assets")

REQUIRED_HEADERS = [
    "AMC_Registration_Number", "Full_Name", "Father_Name", "Profession_or_Degree",
    "Specialty_or_Field", "University_Name", "Graduation_Year", "Exit_Exam_Status",
    "AMC_Registration_Status", "Notes",
]

def cell_text(value):
    return "" if value is None else str(value).strip()

def find_doctor(amc_number):
    if not EXCEL_PATH.exists():
        raise FileNotFoundError("The doctors.xlsx file was not found beside app.py.")
    workbook = load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            try:
                headers = [cell_text(v) for v in next(rows)]
            except StopIteration:
                continue
            if not all(h in headers for h in REQUIRED_HEADERS):
                continue
            indexes = {h: headers.index(h) for h in REQUIRED_HEADERS}
            for row in rows:
                row_values = list(row)
                candidate = cell_text(row_values[indexes["AMC_Registration_Number"]] if indexes["AMC_Registration_Number"] < len(row_values) else "")
                if candidate and candidate.casefold() == amc_number.strip().casefold():
                    return {h: cell_text(row_values[indexes[h]] if indexes[h] < len(row_values) else "") for h in REQUIRED_HEADERS}
        return None
    finally:
        workbook.close()

@app.get("/")
@app.get("/Verify")
@app.get("/verify")
def home():
    return render_template("index.html", query="", result=None, not_found=False, data_error=None)

@app.post("/Verify")
@app.post("/verify")
def verify():
    query = request.form.get("txt_amc_number_search", "").strip()
    if not query:
        return render_template("index.html", query="", result=None, not_found=False, data_error="Please enter an AMC Registration Number."), 400
    try:
        result = find_doctor(query)
        return render_template("index.html", query=query, result=result, not_found=result is None, data_error=None)
    except Exception as exc:
        # Avoid showing server internals; provide a practical local-file message.
        message = "Unable to read the Excel file. Close it in Excel and check that doctors.xlsx is beside app.py."
        app.logger.exception("Excel lookup failed: %s", exc)
        return render_template("index.html", query=query, result=None, not_found=False, data_error=message), 500

if __name__ == "__main__":
    port = int(os.environ.get("PORT", "5000"))
    app.run(host="0.0.0.0", port=port, debug=False)
